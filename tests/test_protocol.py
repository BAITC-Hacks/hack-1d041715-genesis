"""End-to-end demo protocol and DOCX export tests."""

import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from meeting_minutes.export_docx import build_protocol_docx, export_protocol_docx
from meeting_minutes.pipeline import process_meeting


class ProtocolTests(unittest.TestCase):
    """Verify autonomous protocol assembly and downloadable Word output."""

    def setUp(self) -> None:
        self.previous_mode = os.environ.get("APP_MODE")
        os.environ["APP_MODE"] = "demo"

    def tearDown(self) -> None:
        if self.previous_mode is None:
            os.environ.pop("APP_MODE", None)
        else:
            os.environ["APP_MODE"] = self.previous_mode

    def test_demo_pipeline_builds_complete_protocol(self) -> None:
        """Demo mode reaches a structured protocol without an audio path or key."""
        protocol = process_meeting()

        self.assertEqual(protocol.metadata.mode, "demo")
        self.assertGreaterEqual(len(protocol.tasks), 5)
        self.assertGreaterEqual(len(protocol.transcript), 4)
        self.assertTrue(any(segment.speaker_name for segment in protocol.transcript))
        self.assertIn("92", protocol.summary)

    def test_docx_contains_summary_tasks_and_transcript(self) -> None:
        """The generated DOCX includes all required protocol sections."""
        from docx import Document

        protocol = process_meeting()
        content = build_protocol_docx(protocol)
        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = " ".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )

        self.assertTrue(content.startswith(b"PK"))
        self.assertIn("Протокол совещания", text)
        self.assertIn("Демонстрационный режим", text)
        self.assertIn("Нурлан Сагатович", table_text)
        self.assertIn("Поручение", table_text)

    def test_export_writes_nonempty_docx_file(self) -> None:
        """File export creates a valid non-empty Word artifact."""
        protocol = process_meeting()
        with tempfile.TemporaryDirectory() as temp_dir:
            output = export_protocol_docx(protocol, Path(temp_dir) / "protocol.docx")

            self.assertTrue(output.is_file())
            self.assertGreater(output.stat().st_size, 1000)
