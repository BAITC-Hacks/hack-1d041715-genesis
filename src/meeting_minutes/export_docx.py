"""Professional DOCX export for a completed meeting protocol."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from .models import MeetingProtocol


def build_protocol_docx(protocol: MeetingProtocol) -> bytes:
    """Render a protocol to a self-contained DOCX byte string."""
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:
        raise RuntimeError(
            "Не установлена библиотека python-docx. Выполните: pip install -r requirements.txt"
        ) from error

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    _configure_styles(document, Pt, RGBColor, qn)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Протокол совещания")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Источник: {protocol.metadata.source_filename}").italic = True

    if protocol.metadata.mode == "demo":
        notice = document.add_paragraph()
        notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = notice.add_run(
            "Демонстрационный режим: транскрипт и AI-анализ взяты из тестового fixture."
        )
        run.bold = True
    elif protocol.metadata.warnings:
        notice = document.add_paragraph()
        notice.add_run("Ограничение обработки. ").bold = True
        notice.add_run(" ".join(protocol.metadata.warnings))

    _add_heading(document, "Краткое саммари")
    document.add_paragraph(protocol.summary)

    _add_heading(document, "Поручения")
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(0.42), Inches(3.25), Inches(1.55), Inches(1.45))
    headers = ("№", "Поручение", "Ответственный", "Срок")
    for index, (cell, label, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = label
        cell.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
        )
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        _set_cell_fill(cell, "1F4E78", OxmlElement, qn)
        _set_cell_margins(cell, OxmlElement, qn)
    _repeat_table_header(table.rows[0], OxmlElement, qn)

    if protocol.tasks:
        for number, task in enumerate(protocol.tasks, start=1):
            row = table.add_row()
            values = (
                str(number),
                task.task,
                task.assignee or "Не указано",
                task.deadline or "Не указано",
            )
            for index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = width
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell.text = value
                cell.paragraphs[0].alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
                )
                _set_cell_margins(cell, OxmlElement, qn)
                if number % 2 == 0:
                    _set_cell_fill(cell, "EAF2F8", OxmlElement, qn)
    else:
        row = table.add_row()
        row.cells[0].merge(row.cells[-1])
        row.cells[0].text = "Поручения не найдены"
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_margins(row.cells[0], OxmlElement, qn)
    _set_table_borders(table, OxmlElement, qn)

    _add_heading(document, "Транскрипт по говорящим")
    for segment in protocol.transcript:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        label = segment.display_speaker
        timing = f"{_format_time(segment.start)}–{_format_time(segment.end)}"
        paragraph.add_run(f"{label}  {timing}\n").bold = True
        paragraph.add_run(segment.text)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_protocol_docx(protocol: MeetingProtocol, output_path: str | Path) -> Path:
    """Write a protocol DOCX to an explicit path and return the resolved path."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(build_protocol_docx(protocol))
    return path.resolve()


def _configure_styles(document: Any, pt: Any, rgb_color: Any, qn: Any) -> None:
    """Apply a compact formal typography system with Cyrillic-safe fonts."""
    for style_name, size, bold in (
        ("Normal", 11, False),
        ("Title", 22, True),
        ("Heading 1", 15, True),
    ):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = pt(size)
        style.font.bold = bold
        style.font.color.rgb = rgb_color(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    document.styles["Normal"].paragraph_format.space_after = pt(7)
    document.styles["Normal"].paragraph_format.line_spacing = 1.08
    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)


def _add_heading(document: Any, text: str) -> None:
    """Add a black heading kept with the content that follows."""
    paragraph = document.add_paragraph(text, style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = document.styles["Normal"].font.size


def _format_time(seconds: float) -> str:
    """Format a floating-point timestamp as MM:SS."""
    total_seconds = max(0, int(round(seconds)))
    minutes, remainder = divmod(total_seconds, 60)
    return f"{minutes:02d}:{remainder:02d}"


def _set_cell_fill(cell: Any, fill: str, element_factory: Any, qn: Any) -> None:
    """Set a table cell background fill."""
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = element_factory("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, element_factory: Any, qn: Any) -> None:
    """Add consistent breathing room inside a table cell."""
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = element_factory("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", "90"), ("left", "100"), ("bottom", "90"), ("right", "100")):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = element_factory(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table: Any, element_factory: Any, qn: Any) -> None:
    """Apply visible light-gray outer and internal table borders."""
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = element_factory("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = element_factory(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)


def _repeat_table_header(row: Any, element_factory: Any, qn: Any) -> None:
    """Mark the header row to repeat if the task table spans pages."""
    row_properties = row._tr.get_or_add_trPr()
    table_header = element_factory("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    row_properties.append(table_header)
